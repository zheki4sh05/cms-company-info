# -*- coding: utf-8 -*-
"""Generate Word document: functional capabilities by user role (cms-company-info)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

OUTPUT = Path(__file__).resolve().parent.parent / (
    "Перечень_функциональных_возможностей_cms-company-info.docx"
)

ROLES_INTRO = [
    (
        "EXECUTIVE",
        "Исполнительный уровень управления компанией. Полный доступ к изменению "
        "структуры компании, отделов, состава сотрудников и отправке приглашений.",
    ),
    (
        "SUPERVISOR",
        "Руководитель отдела (назначается как manager_id отдела). Используется в "
        "межсервисных сценариях согласования и иерархии подчинения.",
    ),
    (
        "MANAGER",
        "Менеджер / сотрудник с ролью менеджера в компании. Доступ к просмотру "
        "данных компании наравне с другими членами; изменение структуры недоступно.",
    ),
    (
        "Любой сотрудник компании",
        "Пользователь с записью в таблице employee для данной companyId "
        "(роль EXECUTIVE, SUPERVISOR или MANAGER).",
    ),
    (
        "Аутентифицированный пользователь",
        "Пользователь с валидным JWT (Bearer), извлечённым userId; членство в "
        "компании не проверяется.",
    ),
    (
        "Межсервисный вызов",
        "Внутренние HTTP-запросы от других сервисов платформы без проверки роли "
        "сотрудника (доверенная сеть / шлюз).",
    ),
    (
        "Система (Kafka)",
        "Автоматическая обработка событий без участия пользователя.",
    ),
]

FUNCTIONS = [
    # Company
    (
        "Компания",
        "Получение данных текущей компании",
        "Возвращает название компании и число сотрудников по связке userId + employeeId из JWT и заголовка EmployeeId.",
        "Любой сотрудник компании (привязка employee к user)",
        "GET /company",
    ),
    (
        "Компания",
        "Получение идентификатора компании по userId",
        "Служебный запрос: companyId для указанного пользователя.",
        "Межсервисный вызов",
        "GET /company/id/{userId}",
    ),
    (
        "Компания",
        "Изменение названия компании",
        "Обновление поля name первой записи компании в БД.",
        "Межсервисный вызов (в API не требуется JWT и роль)",
        "PATCH /company",
    ),
    (
        "Компания",
        "Удаление компании",
        "Удаление компании и связанных сотрудников (CASCADE).",
        "Межсервисный вызов (в API не требуется JWT и роль)",
        "DELETE /company",
    ),
    # Employee
    (
        "Сотрудники",
        "Получение employeeId текущего пользователя",
        "По userId из JWT возвращается идентификатор записи employee.",
        "Аутентифицированный пользователь",
        "GET /employee/id",
    ),
    (
        "Сотрудники",
        "Список сотрудников компании",
        "Список сотрудников с обогащением профиля из auth-service (email, ФИО, отдел, роль).",
        "Любой сотрудник компании",
        "GET /companies/{companyId}/employees",
    ),
    (
        "Сотрудники",
        "Контекст сотрудника и отдела по userId",
        "employeeId, departmentId, название отдела и роль сотрудника в отделе (departmentRole).",
        "Межсервисный вызов",
        "GET /employee/internal/id/{userId}",
    ),
    (
        "Сотрудники",
        "Руководитель отдела сотрудника",
        "Возвращает SUPERVISOR-руководителя отдела, в котором состоит указанный сотрудник.",
        "Межсервисный вызов",
        "GET /employee/department-manager",
    ),
    (
        "Сотрудники",
        "Подчинённые менеджеры руководителя отдела",
        "Для сотрудника-руководителя отдела (SUPERVISOR как manager_id) — список userId сотрудников с ролью MANAGER в его отделах.",
        "Межсервисный вызов",
        "GET /employee/department-manager-subordinates",
    ),
    # Departments - read
    (
        "Отделы",
        "Список отделов компании",
        "Перечень отделов с количеством сотрудников и данными руководителя (supervisor).",
        "Любой сотрудник компании",
        "GET /companies/{companyId}/departments",
    ),
    (
        "Отделы",
        "Карточка отдела",
        "Детали отдела и список сотрудников отдела (ФИО из auth-service).",
        "Любой сотрудник компании",
        "GET /departments/{id}",
    ),
    # Departments - write EXECUTIVE
    (
        "Отделы",
        "Создание отдела",
        "Создание отдела в компании (название, описание, опционально руководитель).",
        "EXECUTIVE",
        "POST /companies/{companyId}/departments, POST /departments",
    ),
    (
        "Отделы",
        "Изменение отдела",
        "Частичное обновление названия, описания, managerId.",
        "EXECUTIVE",
        "PATCH /companies/{companyId}/departments/{departmentId}, PATCH /departments/{id}",
    ),
    (
        "Отделы",
        "Удаление отдела",
        "Удаление только пустого отдела (без сотрудников в department_employee).",
        "EXECUTIVE",
        "DELETE /companies/{companyId}/departments/{departmentId}, DELETE /departments/{id}",
    ),
    (
        "Отделы",
        "Назначение руководителя отдела (по employeeId)",
        "Установка manager_id и включение руководителя в состав отдела.",
        "EXECUTIVE",
        "POST /departments/{departmentId}/manager",
    ),
    (
        "Отделы",
        "Назначение руководителя отдела (по userId)",
        "Назначение supervisor по userId сотрудника; требуется заголовок companyId.",
        "EXECUTIVE",
        "POST /departments/{departmentId}/supervisor",
    ),
    (
        "Отделы",
        "Добавление сотрудника в отдел",
        "Привязка существующего сотрудника компании к отделу.",
        "EXECUTIVE",
        "POST /companies/{companyId}/departments/{departmentId}/members",
    ),
    (
        "Отделы",
        "Исключение сотрудника из отдела",
        "Удаление связи department_employee; при необходимости сброс manager_id.",
        "EXECUTIVE",
        "DELETE /companies/{companyId}/departments/{departmentId}/members/{employeeId}",
    ),
    (
        "Отделы",
        "Перевод сотрудника между отделами",
        "Перемещение сотрудника из одного отдела в другой в рамках одной компании.",
        "EXECUTIVE",
        "POST /departments/transfer",
    ),
    # Invitations
    (
        "Приглашения",
        "Отправка приглашения в компанию",
        "Создание приглашения и записи employee с ролью MANAGER | SUPERVISOR | EXECUTIVE; "
        "публикация события ADD_USER_COMPANY в Kafka (топик auth).",
        "EXECUTIVE (инициатор; в теле запроса задаётся роль приглашаемого)",
        "POST /invitations/send",
    ),
    # Kafka / system
    (
        "Интеграция",
        "Создание компании по событию Kafka",
        "При событии CREATED в топике company создаётся компания и первый сотрудник с ролью из сообщения (обычно EXECUTIVE).",
        "Система (Kafka)",
        "Consumer: KAFKA_COMPANY_TOPIC, event=CREATED",
    ),
    (
        "Интеграция",
        "Публикация привязки пользователя к компании",
        "После успешного приглашения — сообщение ADD_USER_COMPANY для auth/других сервисов.",
        "Система (инициируется действием EXECUTIVE)",
        "Producer: KAFKA_AUTH_TOPIC",
    ),
    (
        "Интеграция",
        "Взаимодействие с auth-service",
        "Проверка email, получение профилей сотрудников и ФИО для отображения.",
        "Внутренняя логика сервиса (по сценарию вызывающей роли)",
        "HTTP → AUTH_SERVICE_URL",
    ),
    (
        "Документация",
        "Просмотр OpenAPI / Swagger",
        "Интерактивная спецификация REST API сервиса.",
        "Без ограничения роли (доступ к /api-docs)",
        "GET /api-docs, GET /openapi.json",
    ),
]

# Подробное описание: какое «право» что даёт именно в cms-company-info
# (право, где хранится/проверяется, функциональность в сервисе, ограничения)
RIGHTS_IN_SERVICE = [
    {
        "title": "4.1. Роль EXECUTIVE (поле employee.role)",
        "where": (
            "Таблица employee, столбец role. Значение «EXECUTIVE» назначается при создании "
            "первого сотрудника компании (событие Kafka CREATED) или при приглашении "
            "(POST /invitations/send, поле role в теле запроса)."
        ),
        "check": (
            "Middleware requireExecutive и явные проверки member.role !== 'EXECUTIVE' "
            "в invitation.service, department.service (patch, transfer, manager, delete)."
        ),
        "grants": [
            "Создание отделов (POST /companies/{companyId}/departments, POST /departments).",
            "Редактирование и удаление отделов (PATCH, DELETE по id отдела).",
            "Назначение руководителя отдела: POST /departments/{id}/manager и /supervisor.",
            "Добавление и удаление сотрудников в составе отдела (members POST/DELETE).",
            "Перевод сотрудника между отделами (POST /departments/transfer).",
            "Отправка приглашений: POST /invitations/send — только инициатор с ролью EXECUTIVE; "
            "в приглашении можно указать роль приглашаемого: MANAGER, SUPERVISOR или EXECUTIVE.",
            "Публикация в Kafka события ADD_USER_COMPANY после успешного приглашения.",
        ],
        "also_has": [
            "Все операции просмотра, доступные любому сотруднику компании (см. п. 4.4).",
        ],
        "denies": [
            "Нет отдельных «суперправ» вне компании: EXECUTIVE действует только в той companyId, "
            "где пользователь числится сотрудником.",
            "PATCH /company и DELETE /company в текущей реализации не проверяют роль EXECUTIVE "
            "(эндпоинты без JWT) — это не пользовательское право EXECUTIVE, а открытый/служебный API.",
        ],
    },
    {
        "title": "4.2. Роль SUPERVISOR (поле employee.role)",
        "where": (
            "Таблица employee.role. Часто назначается при приглашении или при смене роли "
            "сотрудника. Дополнительно сотрудник может быть закреплён как руководитель отдела "
            "в department.manager_id (см. п. 4.5)."
        ),
        "check": (
            "Для HTTP API пользователя: только requireCompanyMember (роль SUPERVISOR не даёт "
            "прав на изменение). В межсервисных методах: роль SUPERVISOR у manager_id отдела "
            "обязательна в GET /employee/department-manager."
        ),
        "grants": [
            "Просмотр данных компании: GET /company (при наличии employeeId в заголовке).",
            "Список отделов и карточка отдела: GET /companies/{companyId}/departments, GET /departments/{id}.",
            "Список сотрудников компании: GET /companies/{companyId}/employees.",
            "Участие в оргструктуре как руководитель отдела (если назначен manager_id): "
            "другие сервисы получают его через GET /employee/department-manager.",
            "Если назначен manager_id и имеет роль SUPERVISOR: может фигурировать как «руководитель отдела» "
            "в GET /employee/department-manager-subordinates (подчинённые — сотрудники с ролью MANAGER в его отделах).",
        ],
        "also_has": [],
        "denies": [
            "Создание, изменение, удаление отделов; перевод сотрудников; управление составом отдела.",
            "Отправка приглашений (403: Only EXECUTIVE role can perform this action).",
            "Назначение себя или других руководителем через POST /departments/.../manager|supervisor.",
        ],
    },
    {
        "title": "4.3. Роль MANAGER (поле employee.role)",
        "where": (
            "Таблица employee.role. Типичная роль рядового менеджера компании; "
            "назначается при приглашении с role: MANAGER."
        ),
        "check": "requireCompanyMember для операций чтения; запись структуры недоступна.",
        "grants": [
            "Те же операции просмотра, что и у SUPERVISOR: компания, отделы, сотрудники.",
            "GET /employee/id — получение своего employeeId по JWT.",
            "Возврат в GET /employee/internal/id/{userId} поля departmentRole (значение employee.role).",
            "Может быть указан как подчинённый в GET /employee/department-manager-subordinates "
            "(если состоит в отделе, где запрашивающий — руководитель с ролью SUPERVISOR).",
        ],
        "also_has": [],
        "denies": [
            "Любые изменения оргструктуры и приглашения (аналогично SUPERVISOR).",
            "Не используется как валидный «руководитель отдела» в GET /employee/department-manager "
            "(там требуется manager с role = SUPERVISOR).",
        ],
    },
    {
        "title": "4.4. Право «член компании» (членство в employee)",
        "where": (
            "Наличие строки в employee с парой (user_id, company_id). Проверка: "
            "requireCompanyMember / findEmployeeByUserAndCompany."
        ),
        "check": (
            "403 Forbidden: You are not a member of this company — если пользователь "
            "аутентифицирован, но не привязан к companyId из URL или body."
        ),
        "grants": [
            "GET /companies/{companyId}/departments — список отделов своей компании.",
            "GET /companies/{companyId}/employees — список коллег (с данными из auth-service).",
            "GET /departments/{id} — карточка отдела своей компании со списком участников.",
            "GET /company — при совпадении employeeId заголовка и userId из JWT с записью employee.",
        ],
        "also_has": [
            "Само по себе членство не различает EXECUTIVE и MANAGER для чтения — роль влияет только на запись.",
        ],
        "denies": [
            "Доступ к данным чужой companyId (даже с валидным JWT).",
            "Операции изменения без роли EXECUTIVE (см. п. 4.1).",
        ],
    },
    {
        "title": "4.5. Право «руководитель отдела» (department.manager_id)",
        "where": (
            "Не отдельная роль в ACL, а связь: department.manager_id → employee.employee_id. "
            "Назначается только пользователем с правом EXECUTIVE "
            "(POST .../manager, .../supervisor, PATCH отдела с managerId)."
        ),
        "check": (
            "isDepartmentHeadSupervisorInCompany — сотрудник является manager_id хотя бы одного отдела "
            "и имеет employee.role = SUPERVISOR."
        ),
        "grants": [
            "Для других микросервисов: идентификация руководителя сотрудника "
            "(GET /employee/department-manager → userId руководителя с ролью SUPERVISOR).",
            "Для руководителя отдела (SUPERVISOR + manager_id): получение списка userId подчинённых "
            "с ролью MANAGER (GET /employee/department-manager-subordinates).",
            "В списке отделов отображается supervisorId / supervisorName (ФИО из auth-service).",
        ],
        "also_has": [],
        "denies": [
            "Сам по себе факт manager_id не открывает в этом сервисе UI/API для изменения отделов — "
            "только EXECUTIVE может менять структуру.",
            "MANAGER, назначенный manager_id без роли SUPERVISOR, не пройдёт проверку в department-manager API.",
        ],
    },
    {
        "title": "4.6. Аутентификация JWT (Bearer, RS256)",
        "where": (
            "authContextMiddleware: извлечение userId из JWT; опционально employeeId из заголовков "
            "employeeid / x-employee-id."
        ),
        "check": "401 Unauthorized при отсутствии или невалидном токене (requireUser).",
        "grants": [
            "GET /employee/id — employeeId текущего пользователя без проверки companyId.",
            "Базовый доступ к защищённым маршрутам с requireUser (далее — проверка членства/роли).",
            "Проброс Authorization в auth-service при списке сотрудников и приглашении.",
        ],
        "also_has": [
            "Fallback x-user-id для доверенных внутренних вызовов через шлюз.",
        ],
        "denies": [
            "JWT без записи employee не даёт просмотр компании/отделов (нужно членство, п. 4.4).",
            "JWT не заменяет роль EXECUTIVE для операций записи.",
        ],
    },
    {
        "title": "4.7. Роль приглашаемого (поле role в POST /invitations/send)",
        "where": (
            "Тело запроса: role ∈ { MANAGER, SUPERVISOR, EXECUTIVE }. Сохраняется в invitation.role "
            "и employee.role при создании приглашения."
        ),
        "check": (
            "Может задать только EXECUTIVE-инициатор; assertRole в invitation.service."
        ),
        "grants": [
            "MANAGER — приглашённый после принятия получает права п. 4.3 (только просмотр оргструктуры).",
            "SUPERVISOR — права п. 4.2; может быть назначен руководителем отдела EXECUTIVE.",
            "EXECUTIVE — полные права п. 4.1 в рамках компании (в т.ч. приглашать других).",
            "Привязка к departmentId из тела — сотрудник добавляется в указанный отдел.",
            "Событие Kafka ADD_USER_COMPANY синхронизирует привязку userId ↔ companyId с auth/другими сервисами.",
        ],
        "also_has": [],
        "denies": [
            "Нельзя пригласить с ролью вне enum; нельзя пригласить уже существующего в компании сотрудника.",
            "Инициатор приглашения не может быть MANAGER/SUPERVISOR — только EXECUTIVE.",
        ],
    },
    {
        "title": "4.8. Межсервисный доступ (без проверки роли сотрудника)",
        "where": (
            "Эндпоинты без requireUser / requireCompanyMember; вызовы из auth, gateway, "
            "workflow-сервисов в защищённом контуре."
        ),
        "check": "Валидация параметров (userId, employeeId, companyId); без JWT на части маршрутов.",
        "grants": [
            "GET /company/id/{userId} — companyId по пользователю.",
            "GET /employee/internal/id/{userId} — контекст отдела и departmentRole.",
            "GET /employee/department-manager — руководитель SUPERVISOR для сотрудника.",
            "GET /employee/department-manager-subordinates — userId менеджеров под руководителем отдела.",
            "PATCH /company, DELETE /company — изменение/удаление singleton-компании (служебно).",
        ],
        "also_has": [],
        "denies": [
            "Не предназначено для прямого вызова конечным пользователем без контроля шлюза.",
            "Не расширяет права MANAGER/SUPERVISOR на изменение структуры через эти URL.",
        ],
    },
    {
        "title": "4.9. Системное право обработки Kafka (consumer/producer)",
        "where": (
            "Переменные KAFKA_COMPANY_TOPIC, KAFKA_AUTH_TOPIC; запуск consumer при старте приложения."
        ),
        "check": "Обработчик company.consumer — только event === 'CREATED'.",
        "grants": [
            "Вход: CREATED — автоматическое создание company + employee с role из сообщения "
            "(типично первый EXECUTIVE владельца компании).",
            "Выход: ADD_USER_COMPANY после приглашения — уведомление auth о привязке пользователя к companyId.",
        ],
        "also_has": [],
        "denies": [
            "Не обрабатывает произвольные роли пользователей в runtime — только заданные форматы событий.",
            "Не заменяет POST /invitations/send для добавления сотрудников с email-валидацией.",
        ],
    },
]

# Сводная таблица: право → ключевые эндпоинты в cms-company-info
RIGHTS_ENDPOINTS_SUMMARY = [
    ("EXECUTIVE", "Запись оргструктуры и приглашения", "POST/PATCH/DELETE отделов, members, transfer, POST /invitations/send"),
    ("SUPERVISOR", "Чтение + роль в иерархии", "GET company/departments/employees; участие в department-manager*"),
    ("MANAGER", "Только чтение оргструктуры", "GET company, departments, employees, /employee/id"),
    ("Член компании", "Доступ к данным своей companyId", "GET /companies/{companyId}/*, GET /departments/{id}"),
    ("JWT (userId)", "Идентификация пользователя", "GET /employee/id; prerequisite для защищённых маршрутов"),
    ("manager_id отдела", "Руководитель в иерархии (с SUPERVISOR)", "Используется в GET /employee/department-manager*"),
    ("Роль в приглашении", "Права нового сотрудника после onboarding", "employee.role = MANAGER | SUPERVISOR | EXECUTIVE"),
    ("Межсервисный", "Служебные запросы платформы", "GET /company/id/..., /employee/internal/..., PATCH/DELETE /company"),
    ("Kafka", "Автосоздание компании / синхронизация auth", "CREATED in, ADD_USER_COMPANY out"),
]


def add_bullet_list(doc: Document, items: list[str], style: str = "List Bullet") -> None:
    for item in items:
        p = doc.add_paragraph(item, style=style)
        for run in p.runs:
            run.font.size = Pt(11)
            run.font.name = "Times New Roman"


def add_rights_sections(doc: Document) -> None:
    doc.add_heading(
        "4. Права и роли: какая функциональность предоставляется в cms-company-info",
        level=1,
    )
    preamble = doc.add_paragraph(
        "В данном микросервисе «право» — это либо роль сотрудника в компании (employee.role), "
        "либо результат проверки middleware (членство в компании, JWT), либо служебный канал "
        "(межсервисный HTTP, Kafka). Ниже для каждого вида права указано, какие именно операции "
        "и сценарии становятся доступны в рамках cms-company-info, а что явно запрещено кодом сервиса."
    )
    preamble.paragraph_format.first_line_indent = Cm(1.25)

    for block in RIGHTS_IN_SERVICE:
        doc.add_heading(block["title"], level=2)
        for label, key in (
            ("Где задаётся", "where"),
            ("Как проверяется", "check"),
        ):
            p = doc.add_paragraph()
            r = p.add_run(f"{label}: ")
            r.bold = True
            p.add_run(block[key])
            for run in p.runs:
                run.font.size = Pt(11)
                run.font.name = "Times New Roman"

        p = doc.add_paragraph()
        r = p.add_run("Предоставляемая функциональность в сервисе:")
        r.bold = True
        add_bullet_list(doc, block["grants"])

        if block.get("also_has"):
            p = doc.add_paragraph()
            r = p.add_run("Дополнительно:")
            r.bold = True
            add_bullet_list(doc, block["also_has"])

        p = doc.add_paragraph()
        r = p.add_run("Не предоставляет / ограничения:")
        r.bold = True
        add_bullet_list(doc, block["denies"])

        doc.add_paragraph()

    doc.add_heading("4.10. Сводка: право → операции API", level=2)
    add_table(
        doc,
        ["Право / роль", "Назначение в сервисе", "Основные эндпоинты"],
        RIGHTS_ENDPOINTS_SUMMARY,
        [3.5, 5.0, 8.0],
    )


def add_table(doc: Document, headers: list[str], rows: list[tuple], col_widths_cm: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr_cells[i].text = text
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)

    for row_data in rows:
        row = table.add_row().cells
        for i, text in enumerate(row_data):
            row[i].text = str(text)
            for p in row[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)

    for row in table.rows:
        for idx, width in enumerate(col_widths_cm):
            row.cells[idx].width = Cm(width)


def main() -> None:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "Перечень функциональных возможностей\n"
        "микросервиса cms-company-info\n"
        "(TrustFlow)"
    )
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "Times New Roman"

    doc.add_paragraph()
    intro = doc.add_paragraph(
        "Настоящий документ описывает функции, реализованные в сервисе управления "
        "структурой компании (компании, сотрудники, отделы, приглашения), с указанием "
        "роли пользователя или типа доступа, которому доступна соответствующая операция. "
        "Сведения сформированы по фактической реализации REST API, middleware проверки "
        "доступа и интеграций (PostgreSQL, Kafka, auth-service)."
    )
    intro.paragraph_format.first_line_indent = Cm(1.25)

    doc.add_heading("1. Роли и типы доступа в системе", level=1)
    add_table(
        doc,
        ["Обозначение", "Описание"],
        ROLES_INTRO,
        [4.5, 12.0],
    )

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Аутентификация: для пользовательских операций используется JWT (RS256) в заголовке "
        "Authorization: Bearer. Идентификатор сотрудника передаётся в заголовках employeeid / "
        "x-employee-id. Для доверенных внутренних вызовов допускается передача x-user-id."
    )
    note.paragraph_format.first_line_indent = Cm(1.25)

    add_rights_sections(doc)

    doc.add_heading("5. Перечень функциональных возможностей", level=1)
    func_rows = []
    for n, (module, name, desc, role, api) in enumerate(FUNCTIONS, start=1):
        func_rows.append((n, module, name, desc, role, api))

    add_table(
        doc,
        [
            "№",
            "Модуль",
            "Функция",
            "Описание",
            "Доступ (роль / тип)",
            "API / канал",
        ],
        func_rows,
        [1.0, 2.5, 3.5, 5.5, 3.5, 4.0],
    )

    doc.add_heading("6. Матрица доступа по ролям сотрудника компании", level=1)
    matrix = [
        ("Просмотр компании, отделов, списка сотрудников", "Да", "Да", "Да"),
        ("Создание / изменение / удаление отделов", "Да", "Нет", "Нет"),
        ("Управление составом отделов и перевод между отделами", "Да", "Нет", "Нет"),
        ("Назначение руководителя отдела (SUPERVISOR)", "Да", "Нет", "Нет"),
        ("Отправка приглашений в компанию", "Да", "Нет", "Нет"),
    ]
    add_table(
        doc,
        ["Операция", "EXECUTIVE", "SUPERVISOR", "MANAGER"],
        matrix,
        [8.0, 2.5, 2.5, 2.5],
    )

    doc.add_paragraph()
    footer = doc.add_paragraph(
        "Дата формирования: май 2026 г. Источник: репозиторий cms-company-info, "
        "OpenAPI (src/docs/openapi.json), middleware requireExecutive / requireCompanyMember."
    )
    footer.paragraph_format.space_before = Pt(12)

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
